"""Predição em lote e por XLSX (#15).

O lote é tudo ou nada: todas as solicitações passam pelas mesmas validações
da predição individual (`planejar`) antes de qualquer previsão ou gravação,
e os problemas voltam juntos, cada um com seu campo (e a linha, na planilha).
"""

import re
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pydantic import ValidationError

from case_mr_health.api.erros import ErroApi, LoteInvalido, _detalhe
from case_mr_health.api.historico import Historico
from case_mr_health.api.modelos import CatalogoModelos
from case_mr_health.api.schemas.erros import DetalheErro
from case_mr_health.api.schemas.lote import (
    LIMITE_LOTE,
    ItemLote,
    RespostaLote,
)
from case_mr_health.api.schemas.predicao import PedidoPredicao
from case_mr_health.api.servicos import predicao
from case_mr_health.config import Settings
from case_mr_health.particoes import DIAS_POR_PERIODO
from case_mr_health.treino import horizonte_maximo

ABA = 'Solicitações'
ABA_INSTRUCOES = 'Instruções'
COLUNAS = ('sku', 'granularidade', 'data_referencia', 'horizonte')
OBRIGATORIAS = ('sku', 'granularidade')
MIME_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


@dataclass(frozen=True)
class Solicitacao:
    indice: int
    linha: int | None
    pedido: PedidoPredicao


def _campo(s: Solicitacao, campo: str | None) -> str:
    indice, linha = s.indice, s.linha
    if linha is not None:
        return f'linha {linha}' + (f' / {campo}' if campo else '')
    return f'itens.{indice}' + (f'.{campo}' if campo else '')


def prever_lote(
    solicitacoes: list[Solicitacao],
    modelos: CatalogoModelos,
    historico: Historico,
    cfg: Settings,
    problemas_previos: list[DetalheErro] | None = None,
) -> RespostaLote:
    """`problemas_previos`: erros de leitura da planilha, já com a linha."""
    planos, problemas = [], list(problemas_previos or [])
    for s in solicitacoes:
        try:
            planos.append(predicao.planejar(s.pedido, modelos, cfg))
        except ErroApi as erro:
            problemas.append(
                DetalheErro(
                    campo=_campo(s, erro.campo), mensagem=erro.mensagem
                )
            )
    if problemas:
        raise LoteInvalido(
            f'{len(problemas)} problema(s) no lote; nada foi previsto nem '
            'gravado.',
            detalhes=sorted(problemas, key=_ordem),
        )
    resultados = [
        ItemLote(
            indice=s.indice,
            linha=s.linha,
            previsao=predicao.executar(plano, historico, cfg),
        )
        for s, plano in zip(solicitacoes, planos, strict=True)
    ]
    return RespostaLote(total=len(resultados), resultados=resultados)


def _ordem(detalhe: DetalheErro) -> int:
    achada = re.match(r'linha (\d+)', detalhe.campo)
    return int(achada.group(1)) if achada else 0


def solicitacoes_do_json(pedidos: list[PedidoPredicao]) -> list[Solicitacao]:
    return [Solicitacao(i, None, p) for i, p in enumerate(pedidos)]


# --- XLSX -----------------------------------------------------------------


def _celula(valor):
    """Valor de célula no formato que o contrato aceita."""
    if isinstance(valor, str):
        valor = valor.strip() or None
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)
    return valor


def _abrir(conteudo: bytes):
    try:
        livro = load_workbook(
            BytesIO(conteudo), read_only=True, data_only=True
        )
    except (BadZipFile, KeyError, OSError, ValueError) as erro:
        raise LoteInvalido(
            'O arquivo não é uma planilha XLSX válida.',
            detalhes=[
                DetalheErro(
                    campo='arquivo', mensagem='não foi possível ler o XLSX'
                )
            ],
        ) from erro
    return livro[ABA] if ABA in livro.sheetnames else livro.worksheets[0]


def _cabecalho(linhas) -> dict[str, int]:
    primeira = next(linhas, ())
    nomes = {
        str(c).strip().lower(): i
        for i, c in enumerate(primeira)
        if c is not None
    }
    faltando = [c for c in OBRIGATORIAS if c not in nomes]
    if faltando:
        raise LoteInvalido(
            'O cabeçalho da planilha não tem as colunas obrigatórias.',
            detalhes=[
                DetalheErro(
                    campo=f'linha 1 / {coluna}',
                    mensagem='coluna obrigatória ausente no cabeçalho',
                )
                for coluna in faltando
            ],
        )
    return nomes


def solicitacoes_do_xlsx(
    conteudo: bytes,
) -> tuple[list[Solicitacao], list[DetalheErro]]:
    """Lê a planilha: solicitações válidas e problemas de leitura por linha.

    Erros estruturais (arquivo, cabeçalho, limite, planilha vazia) levantam
    `LoteInvalido`; os de linha voltam para somar-se aos de regra de negócio.
    """
    folha = _abrir(conteudo)
    linhas = folha.iter_rows(values_only=True)
    colunas = _cabecalho(linhas)
    solicitacoes, problemas = [], []
    for numero, valores in enumerate(linhas, start=2):
        bruto = {
            c: _celula(valores[i] if i < len(valores) else None)
            for c, i in colunas.items()
            if c in COLUNAS
        }
        if all(v is None for v in bruto.values()):
            continue
        if len(solicitacoes) + len(problemas) >= LIMITE_LOTE:
            raise LoteInvalido(
                f'A planilha passa de {LIMITE_LOTE} solicitações.',
                detalhes=[
                    DetalheErro(
                        campo=f'linha {numero}',
                        mensagem=f'limite de {LIMITE_LOTE} solicitações',
                    )
                ],
            )
        try:
            pedido = PedidoPredicao.model_validate({
                c: v for c, v in bruto.items() if v is not None
            })
        except ValidationError as erro:
            problemas += [
                DetalheErro(
                    campo=f'linha {numero} / {(d := _detalhe(e)).campo}',
                    mensagem=d.mensagem,
                )
                for e in erro.errors()
            ]
            continue
        solicitacoes.append(
            Solicitacao(len(solicitacoes) + len(problemas), numero, pedido)
        )
    if not solicitacoes and not problemas:
        raise LoteInvalido(
            'A planilha não tem solicitações.',
            detalhes=[
                DetalheErro(
                    campo='arquivo', mensagem='nenhuma linha preenchida'
                )
            ],
        )
    return solicitacoes, problemas


def template_xlsx(modelos: CatalogoModelos, cfg: Settings) -> bytes:
    """Template com colunas, exemplos válidos e instruções de preenchimento."""
    livro = Workbook()
    aba = livro.active
    aba.title = ABA
    aba.append(list(COLUNAS))
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for sku, granularidade in _exemplos(modelos):
        aba.append([sku, granularidade, None, 1])
    aba.freeze_panes = 'A2'
    for coluna, largura in zip('ABCD', (14, 16, 18, 12), strict=True):
        aba.column_dimensions[coluna].width = largura
    instrucoes = livro.create_sheet(ABA_INSTRUCOES)
    for linha in _instrucoes(cfg):
        instrucoes.append([linha])
    instrucoes.column_dimensions['A'].width = 100
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _exemplos(modelos: CatalogoModelos) -> list[tuple[str, str]]:
    """Uma linha por granularidade, sem data: vale o fim do treino."""
    exemplos = {}
    for sku, granularidade in sorted(modelos.modelos):
        exemplos.setdefault(granularidade, (sku, granularidade))
    if exemplos:
        return list(exemplos.values())
    return [('item A', 'diaria'), ('item A', 'semanal')]


def _instrucoes(cfg: Settings) -> list[str]:
    diaria = horizonte_maximo(cfg, 'diaria')
    semanal = horizonte_maximo(cfg, 'semanal')
    return [
        'Preencha uma solicitação por linha da aba "Solicitações".',
        '',
        'sku (obrigatório): `item A` ou só `A`.',
        'granularidade (obrigatório): `diaria` ou `semanal`. A mensal não é '
        'servida.',
        'data_referencia (opcional): último período conhecido, no formato '
        'AAAA-MM-DD ou como data do Excel. Em branco, vale o fim do treino '
        'do modelo servido. Na semanal, vale a segunda-feira da semana. Não '
        'pode ser anterior ao fim do treino do modelo.',
        f'horizonte (opcional, padrão 1): períodos a prever, inteiro maior '
        f'que zero. Somado ao intervalo entre o fim do treino e a data de '
        f'referência, não pode passar de {diaria} dias ou {semanal} semanas.',
        '',
        f'Máximo de {LIMITE_LOTE} solicitações por arquivo. Linhas em '
        'branco são ignoradas.',
        'Qualquer erro cancela o arquivo inteiro; a resposta indica a '
        'linha e a coluna de cada problema.',
        f'Dias por período: diária {DIAS_POR_PERIODO["diaria"]}, semanal '
        f'{DIAS_POR_PERIODO["semanal"]}.',
        'As previsões são da rede, não de uma unidade.',
    ]
