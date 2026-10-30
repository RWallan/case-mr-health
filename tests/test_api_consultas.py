from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace

import duckdb
import pytest
from faker import Faker
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from case_mr_health.api.erros import registrar_tratadores
from case_mr_health.api.historico import Historico, LinhaHistorico
from case_mr_health.api.rotas import consultas as rotas
from case_mr_health.api.schemas.consultas import VersaoDisponivel
from case_mr_health.api.servicos import consultas
from case_mr_health.config import Settings

fake = Faker('pt_BR')
HTTP_OK = 200
HTTP_SERVICE_UNAVAILABLE = 503


@pytest.fixture
def configuracao(tmp_path) -> Settings:
    return Settings(
        _env_file=None,
        historico_db_path=tmp_path / 'historico.duckdb',
        db_path=tmp_path / 'base.duckdb',
        mlflow_tracking_uri=f'sqlite:///{tmp_path / "mlflow.db"}',
        mlflow_artifact_dir=tmp_path / 'artefatos',
    )


@pytest.fixture
def app(configuracao):
    app = FastAPI()
    app.include_router(rotas.router)
    registrar_tratadores(app)
    app.state.cfg = configuracao
    app.state.historico = Historico(configuracao.historico_db_path)
    yield app
    app.state.historico.fechar()


@pytest.fixture
def cliente(app):
    with TestClient(app) as cliente:
        yield cliente


def _linha(indice: int = 0) -> LinhaHistorico:
    fake.seed_instance(indice + 17)
    alvo = date(2021, 8, 20) + timedelta(days=indice)
    quantidade = float(fake.random_int(min=1, max=40))
    valor_unitario = float(fake.random_int(min=5, max=25))
    return LinhaHistorico(
        id_requisicao=fake.uuid4(),
        sku=fake.random_element(('item A', 'item B')),
        granularidade=fake.random_element(('diaria', 'semanal')),
        data_geracao=datetime.now(UTC).replace(tzinfo=None),
        data_referencia=alvo - timedelta(days=1),
        data_alvo=alvo,
        versao_modelo=fake.numerify(text='##'),
        hash_artefato=fake.sha256(),
        quantidade_prevista=quantidade,
        valor_unitario=valor_unitario,
        valor_previsto=quantidade * valor_unitario,
        limite_inferior_80=0.0,
        limite_superior_80=quantidade + fake.random_int(min=1, max=5),
        limite_inferior_95=0.0,
        limite_superior_95=quantidade + fake.random_int(min=6, max=10),
        experimento_id=fake.numerify(text='###'),
        run_id=fake.uuid4().replace('-', ''),
    )


def test_historico_tem_paginacao_e_filtros(cliente, app):
    linhas = [_linha(indice) for indice in range(7)]
    app.state.historico.registrar(linhas)
    esperadas = [
        linha
        for linha in linhas
        if linha.sku == linhas[0].sku
        and linha.granularidade == linhas[0].granularidade
    ]

    resposta = cliente.get(
        '/previsoes',
        params={
            'sku': linhas[0].sku,
            'granularidade': linhas[0].granularidade,
            'limite': 1,
        },
    )

    assert resposta.status_code == HTTP_OK
    corpo = resposta.json()
    assert corpo['total'] == len(esperadas)
    assert corpo['limite'] == 1
    assert corpo['deslocamento'] == 0
    assert corpo['itens'][0]['sku'] == linhas[0].sku


def test_exportacao_xlsx_contem_colunas_e_previsoes(cliente, app):
    linha = _linha()
    app.state.historico.registrar([linha])

    resposta = cliente.get('/previsoes/exportacao.xlsx')

    assert resposta.status_code == HTTP_OK
    assert resposta.headers['content-type'].startswith(
        'application/vnd.openxmlformats-officedocument'
    )
    livro = load_workbook(BytesIO(resposta.content), read_only=True)
    aba = livro['Previsões']
    assert tuple(next(aba.iter_rows(values_only=True))) == (
        'SKU',
        'Granularidade',
        'Período-alvo',
        'Quantidade prevista',
        'Valor unitário',
        'Valor previsto',
        'Limite inferior 80%',
        'Limite superior 80%',
        'Limite inferior 95%',
        'Limite superior 95%',
        'Recomendação',
        'Versão do modelo',
        'Hash do artefato',
        'Data de geração',
        'ID da requisição',
    )
    dados = next(aba.iter_rows(min_row=2, values_only=True))
    assert dados[0] == linha.sku
    assert dados[2] == datetime.combine(linha.data_alvo, datetime.min.time())
    assert dados[3] == linha.quantidade_prevista


def test_exportacao_traz_a_recomendacao_so_onde_ela_existe(cliente, app):
    com = replace(_linha(3), recomendacao='aumento_significativo')
    sem = replace(_linha(4), recomendacao=None)
    app.state.historico.registrar([com, sem])

    resposta = cliente.get('/previsoes/exportacao.xlsx')

    aba = load_workbook(BytesIO(resposta.content), read_only=True)['Previsões']
    cabecalho, *linhas = list(aba.iter_rows(values_only=True))
    indice = cabecalho.index('Recomendação')
    por_sku_data = {(r[0], r[2].date()): r[indice] for r in linhas}
    assert por_sku_data[com.sku, com.data_alvo] == 'aumento_significativo'
    assert por_sku_data[sem.sku, sem.data_alvo] is None


def test_historico_antigo_ganha_a_coluna_de_recomendacao(tmp_path):
    caminho = tmp_path / 'antigo.duckdb'
    with duckdb.connect(str(caminho)) as con:
        con.execute(
            'CREATE TABLE previsoes (id_requisicao VARCHAR NOT NULL, '
            'sku VARCHAR NOT NULL, granularidade VARCHAR NOT NULL, '
            'data_geracao TIMESTAMP NOT NULL, data_referencia DATE NOT NULL, '
            'data_alvo DATE NOT NULL, versao_modelo VARCHAR NOT NULL, '
            'hash_artefato VARCHAR NOT NULL, '
            'quantidade_prevista DOUBLE NOT NULL, '
            'valor_unitario DOUBLE NOT NULL, valor_previsto DOUBLE NOT NULL, '
            'limite_inferior_80 DOUBLE, limite_superior_80 DOUBLE, '
            'limite_inferior_95 DOUBLE, limite_superior_95 DOUBLE, '
            'experimento_id VARCHAR NOT NULL, run_id VARCHAR NOT NULL, '
            'PRIMARY KEY (id_requisicao, data_alvo))'
        )
    historico = Historico(caminho)
    try:
        historico.registrar([replace(_linha(5), recomendacao='estavel')])
        linhas, total = historico.consultar(None, None)
    finally:
        historico.fechar()

    assert total == 1
    assert linhas[0].recomendacao == 'estavel'


def test_exportacao_respeita_filtros(cliente, app):
    primeira, segunda = _linha(1), _linha(2)
    segunda = replace(segunda, sku=primeira.sku, granularidade='diaria')
    app.state.historico.registrar([primeira, segunda])

    resposta = cliente.get(
        '/previsoes/exportacao.xlsx',
        params={'sku': primeira.sku, 'granularidade': 'diaria'},
    )

    assert resposta.status_code == HTTP_OK
    aba = load_workbook(BytesIO(resposta.content), read_only=True)['Previsões']
    assert sum(1 for _ in aba.iter_rows(min_row=2)) == 1


def test_modelos_disponiveis_sao_devolvidos(cliente, monkeypatch):
    esperados = [
        VersaoDisponivel(
            sku='item A',
            granularidade='diaria',
            versao='1',
            aliases=['champion'],
            champion=True,
            algoritmo='naive',
            hash_artefato=fake.sha256(),
            run_id=fake.uuid4().replace('-', ''),
            criada_em=datetime.now(UTC),
        )
    ]
    monkeypatch.setattr(consultas, 'listar_modelos', lambda cfg: esperados)

    resposta = cliente.get('/modelos')

    assert resposta.status_code == HTTP_OK
    assert resposta.json()['modelos'][0]['champion'] is True


def test_listar_modelos_retorna_alias_e_rastreabilidade(
    configuracao, monkeypatch
):
    versao = SimpleNamespace(
        version='4',
        tags={'algoritmo': 'tsb', 'hash_artefato': fake.sha256()},
        run_id=fake.uuid4().replace('-', ''),
        creation_timestamp=1_798_000_000_000,
    )

    def buscar_versoes(filtro):
        if filtro == "name = 'demanda-item-a-diaria'":
            return [versao]
        return []

    cliente = SimpleNamespace(
        search_model_versions=buscar_versoes,
        get_registered_model=lambda nome: SimpleNamespace(
            aliases={'champion': 4, 'selecionado': 3}
        ),
    )
    monkeypatch.setattr(consultas, 'conectar', lambda cfg: cliente)

    modelos = consultas.listar_modelos(configuracao)

    assert len(modelos) == 1
    assert modelos[0].sku == 'item A'
    assert modelos[0].aliases == ['champion']
    assert modelos[0].champion is True
    assert modelos[0].algoritmo == 'tsb'
    assert modelos[0].hash_artefato == versao.tags['hash_artefato']
    assert modelos[0].run_id == versao.run_id


def test_serie_diaria_le_a_base_e_marca_zeros(cliente, configuracao):
    rng = fake
    inicio = date(2021, 6, 1)
    valores = [rng.random_int(min=0, max=5) for _ in range(5)]
    imputados = [rng.boolean() for _ in valores]
    con = duckdb.connect(str(configuracao.db_path))
    con.execute(
        'CREATE TABLE serie_diaria '
        '(sku VARCHAR, data DATE, quantidade BIGINT, imputado BOOLEAN)'
    )
    con.executemany(
        'INSERT INTO serie_diaria VALUES (?, ?, ?, ?)',
        [
            ('item A', inicio + timedelta(days=i), valor, imputado)
            for i, (valor, imputado) in enumerate(
                zip(valores, imputados, strict=True)
            )
        ],
    )
    con.close()

    resposta = cliente.get('/series/A/diaria')

    assert resposta.status_code == HTTP_OK
    pontos = resposta.json()['pontos']
    assert [ponto['quantidade'] for ponto in pontos] == valores
    assert [ponto['imputado'] for ponto in pontos] == imputados
    assert pontos[0]['data'] == inicio.isoformat()


def test_serie_semanal_marca_imputado_como_falso(cliente, configuracao):
    inicio = date(2021, 6, 7)
    valores = [fake.random_int(min=0, max=50) for _ in range(4)]
    con = duckdb.connect(str(configuracao.db_path))
    con.execute(
        'CREATE TABLE serie_semanal '
        '(sku VARCHAR, semana DATE, quantidade BIGINT)'
    )
    con.executemany(
        'INSERT INTO serie_semanal VALUES (?, ?, ?)',
        [
            ('item A', inicio + timedelta(days=7 * i), valor)
            for i, valor in enumerate(valores)
        ],
    )
    con.close()

    resposta = cliente.get('/series/item%20A/semanal')

    assert resposta.status_code == HTTP_OK
    pontos = resposta.json()['pontos']
    assert [ponto['quantidade'] for ponto in pontos] == valores
    assert all(not ponto['imputado'] for ponto in pontos)


def test_serie_sem_base_responde_503(cliente):
    resposta = cliente.get('/series/A/diaria')

    assert resposta.status_code == HTTP_SERVICE_UNAVAILABLE
    assert resposta.json()['codigo'] == 'historico_indisponivel'


def test_modelos_documenta_catalogo_e_exportacao_no_openapi(cliente):
    esquema = cliente.get('/openapi.json').json()

    assert '/modelos' in esquema['paths']
    assert '/previsoes/exportacao.xlsx' in esquema['paths']
    assert '/series/{sku}/{granularidade}' in esquema['paths']
    conteudo = esquema['paths']['/previsoes/exportacao.xlsx']['get'][
        'responses'
    ]['200']['content']
    assert any(tipo.endswith('spreadsheetml.sheet') for tipo in conteudo)
