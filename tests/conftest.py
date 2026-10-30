from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import pytest
from faker import Faker
from openpyxl import Workbook

from case_mr_health.config import Settings


@dataclass
class Planilhas:
    """Planilhas sintéticas e os valores esperados, calculados à parte."""

    raw_dir: Path
    pedidos: list[tuple[int, str]]
    itens: list[tuple[int, str, int]]
    precos: dict[str, int]

    @property
    def quantidade_total(self) -> int:
        return sum(q for _, _, q in self.itens)

    @property
    def valor_por_pedido(self) -> dict[int, int]:
        valores: dict[int, int] = defaultdict(int)
        for id_pedido, sku, qtd in self.itens:
            valores[id_pedido] += qtd * self.precos[sku]
        return dict(valores)

    @property
    def quantidade_por_par(self) -> dict[tuple[int, str], int]:
        pares: dict[tuple[int, str], int] = defaultdict(int)
        for id_pedido, sku, qtd in self.itens:
            pares[id_pedido, sku] += qtd
        return dict(pares)

    @property
    def pares_repetidos(self) -> set[tuple[int, str]]:
        contagem: dict[tuple[int, str], int] = defaultdict(int)
        for id_pedido, sku, _ in self.itens:
            contagem[id_pedido, sku] += 1
        return {par for par, n in contagem.items() if n > 1}


def escrever_planilhas(
    raw_dir: Path,
    pedidos: list[tuple],
    itens: list[tuple],
    precos: dict[str, int],
) -> None:
    """Replica o layout das planilhas reais (coluna de índice sem nome)."""
    raw_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.append([None, 'ID_PEDIDO', 'DATA', 'VALOR_TOTAL'])
    for i, (id_pedido, data) in enumerate(pedidos):
        ws.append([i, id_pedido, data, None])
    wb.save(raw_dir / 'PEDIDO-_1_.xlsx')

    wb = Workbook()
    ws = wb.active
    ws.append([None, 'ID_PEDIDO', 'ID_ITEM', 'QUANTIDADE'])
    for i, (id_pedido, sku, qtd) in enumerate(itens):
        ws.append([i, id_pedido, sku, qtd])
    wb.save(raw_dir / 'ITEM_PEDIDO-_2_.xlsx')

    wb = Workbook()
    ws = wb.active
    ws.append([None, 0])
    for sku, preco in precos.items():
        ws.append([sku, preco])
    wb.save(raw_dir / 'ITENS-_3_.xlsx')


@pytest.fixture
def faker_br() -> Faker:
    fake = Faker('pt_BR')
    Faker.seed(20211)
    return fake


@pytest.fixture
def planilhas(tmp_path, faker_br) -> Planilhas:
    precos = {f'item {letra}': faker_br.random_int(5, 40) for letra in 'ABCD'}
    pedidos = []
    for seq in range(1, 121):
        dia = faker_br.date_between(date(2021, 6, 1), date(2021, 9, 1))
        pedidos.append((int(f'{seq}{dia:%Y%m%d}'), f'{dia:%Y-%m-%d}'))
    itens = []
    for id_pedido, _ in pedidos:
        for sku in faker_br.random_elements(
            list(precos), length=faker_br.random_int(1, 3), unique=True
        ):
            itens.append((id_pedido, sku, faker_br.random_int(1, 5)))
    # Dois pares (pedido, SKU) repetidos, como nas planilhas reais.
    for id_pedido, sku, _ in (itens[0], itens[5]):
        itens.append((id_pedido, sku, faker_br.random_int(1, 5)))
    raw_dir = tmp_path / 'raw'
    escrever_planilhas(raw_dir, pedidos, itens, precos)
    return Planilhas(raw_dir, pedidos, itens, precos)


@pytest.fixture
def settings(tmp_path, planilhas) -> Settings:
    return Settings(
        raw_dir=planilhas.raw_dir,
        db_path=tmp_path / 'saida' / 'base.duckdb',
        relatorio_path=tmp_path / 'saida' / 'relatorio.md',
        desenvolvimento_db_path=tmp_path / 'saida' / 'desenvolvimento.duckdb',
        oot_db_path=tmp_path / 'saida' / 'oot.duckdb',
    )
