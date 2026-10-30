"""Predição em lote e por XLSX (#15): Seam 3 com TestClient.

Reaproveita os modelos pré-registrados de `test_api_predicao`: `item A` com
as duas granularidades, `item B` diária e semanal, `item D` só semanal, e
`item C` sem nenhum modelo servível.
"""

from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from case_mr_health.api.schemas.lote import LIMITE_LOTE
from tests.test_api_predicao import (  # noqa: F401
    FIM_DIARIA,
    FIM_SEMANAL,
    base,
    cliente,
    registro,
)

HTTP_OK = 200
HTTP_INVALIDO = 422
EXEMPLOS_DO_TEMPLATE = 2  # uma linha por granularidade
COLUNAS = ('sku', 'granularidade', 'data_referencia', 'horizonte')


def _item(sku='item A', granularidade='diaria', **campos) -> dict:
    fim = FIM_DIARIA if granularidade == 'diaria' else FIM_SEMANAL
    return {
        'sku': sku,
        'granularidade': granularidade,
        'data_referencia': str(fim),
        **campos,
    }


def _planilha(linhas: list[list], cabecalho=COLUNAS) -> bytes:
    wb = Workbook()
    wb.active.title = 'Solicitações'
    wb.active.append(list(cabecalho))
    for linha in linhas:
        wb.active.append(linha)
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _enviar(cliente, conteudo: bytes):
    return cliente.post(
        '/predicoes/lote/xlsx', files={'arquivo': ('lote.xlsx', conteudo)}
    )


def _campos(resposta) -> dict[str, str]:
    return {d['campo']: d['mensagem'] for d in resposta.json()['detalhes']}


def _gravadas(cliente) -> int:
    return cliente.app.state.historico.consultar(None, None, limite=None)[1]


# --- lote em JSON -------------------------------------------------------


def test_lote_preve_cada_item_na_ordem_enviada(cliente):
    itens = [
        _item(horizonte=2),
        _item('item A', 'semanal'),
        _item('item B', 'diaria', horizonte=3),
    ]

    resposta = cliente.post('/predicoes/lote', json={'itens': itens})

    assert resposta.status_code == HTTP_OK
    corpo = resposta.json()
    assert corpo['total'] == len(itens)
    assert [r['indice'] for r in corpo['resultados']] == [0, 1, 2]
    previsoes = [r['previsao'] for r in corpo['resultados']]
    assert [(p['sku'], p['granularidade']) for p in previsoes] == [
        ('item A', 'diaria'),
        ('item A', 'semanal'),
        ('item B', 'diaria'),
    ]
    assert [len(p['periodos']) for p in previsoes] == [2, 1, 3]
    assert all(r['linha'] is None for r in corpo['resultados'])


def test_item_do_lote_igual_a_predicao_individual(cliente):
    item = _item('item B', 'diaria', horizonte=3)

    individual = cliente.post('/predicoes', json=item).json()
    em_lote = cliente.post('/predicoes/lote', json={'itens': [item]}).json()

    previsao = em_lote['resultados'][0]['previsao']
    assert [p['quantidade_prevista'] for p in previsao['periodos']] == [
        p['quantidade_prevista'] for p in individual['periodos']
    ]
    assert previsao['versao_modelo'] == individual['versao_modelo']
    assert previsao['hash_artefato'] == individual['hash_artefato']


def test_lote_grava_cada_periodo_no_historico(cliente):
    antes = _gravadas(cliente)

    cliente.post(
        '/predicoes/lote',
        json={'itens': [_item(horizonte=2), _item('item A', 'semanal')]},
    )

    assert _gravadas(cliente) == antes + 3


def test_lote_invalido_lista_todos_os_problemas_e_nao_grava_nada(cliente):
    antes = _gravadas(cliente)
    itens = [
        _item('item Z'),
        _item(horizonte=15),
        _item('item C'),
        _item(),
    ]

    resposta = cliente.post('/predicoes/lote', json={'itens': itens})

    assert resposta.status_code == HTTP_INVALIDO
    assert resposta.json()['codigo'] == 'lote_invalido'
    campos = _campos(resposta)
    assert set(campos) == {'itens.0.sku', 'itens.1.horizonte', 'itens.2'}
    assert 'SKU desconhecido' in campos['itens.0.sku']
    assert 'Sem modelo para item C/diaria' in campos['itens.2']
    assert _gravadas(cliente) == antes


def test_lote_json_aceita_a_data_ausente(cliente):
    itens = [
        {'sku': 'item A', 'granularidade': 'diaria', 'horizonte': 2},
        {'sku': 'item A', 'granularidade': 'semanal'},
    ]

    resposta = cliente.post('/predicoes/lote', json={'itens': itens})

    assert resposta.status_code == HTTP_OK
    referencias = [
        r['previsao']['data_referencia'] for r in resposta.json()['resultados']
    ]
    assert referencias == [str(FIM_DIARIA), str(FIM_SEMANAL - timedelta(6))]


def test_lote_aplica_as_mesmas_regras_da_individual(cliente):
    item = _item(data_referencia='2021-08-01')

    individual = cliente.post('/predicoes', json=item)
    em_lote = cliente.post('/predicoes/lote', json={'itens': [item]})

    assert individual.status_code == HTTP_INVALIDO
    assert em_lote.status_code == HTTP_INVALIDO
    assert (
        _campos(em_lote)['itens.0.data_referencia']
        == (_campos(individual)['data_referencia'])
    )


def test_erro_de_contrato_indica_o_item_e_o_campo(cliente):
    itens = [_item(), _item(granularidade='mensal'), {'sku': 'A'}]

    resposta = cliente.post('/predicoes/lote', json={'itens': itens})

    assert resposta.status_code == HTTP_INVALIDO
    campos = _campos(resposta)
    assert 'mensal não é servida' in campos['itens.1.granularidade']
    assert campos['itens.2.granularidade'] == 'campo obrigatório ausente'
    assert 'itens.2.data_referencia' not in campos


@pytest.mark.parametrize(
    ('itens', 'campo'),
    [([], 'itens'), ([_item()] * (LIMITE_LOTE + 1), 'itens')],
)
def test_lote_vazio_ou_grande_demais_e_rejeitado(cliente, itens, campo):
    resposta = cliente.post('/predicoes/lote', json={'itens': itens})

    assert resposta.status_code == HTTP_INVALIDO
    assert campo in _campos(resposta)


# --- template ---------------------------------------------------------------


def test_template_tem_colunas_exemplos_e_instrucoes(cliente):
    resposta = cliente.get('/predicoes/lote/template.xlsx')

    assert resposta.status_code == HTTP_OK
    assert 'spreadsheetml' in resposta.headers['content-type']
    assert 'attachment' in resposta.headers['content-disposition']
    livro = load_workbook(BytesIO(resposta.content))
    assert livro.sheetnames == ['Solicitações', 'Instruções']
    linhas = list(livro['Solicitações'].iter_rows(values_only=True))
    assert linhas[0] == COLUNAS
    exemplos = linhas[1:]
    assert {e[1] for e in exemplos} == {'diaria', 'semanal'}
    instrucoes = ' '.join(
        str(c[0].value) for c in livro['Instruções'].iter_rows()
    )
    assert 'data_referencia (opcional)' in instrucoes
    assert all(e[2] is None for e in exemplos)
    assert '14 dias' in instrucoes
    assert '2 semanas' in instrucoes


def test_template_preenchido_com_os_exemplos_e_aceito(cliente):
    template = cliente.get('/predicoes/lote/template.xlsx').content

    resposta = _enviar(cliente, template)

    assert resposta.status_code == HTTP_OK
    assert resposta.json()['total'] == EXEMPLOS_DO_TEMPLATE


# --- lote em XLSX --------------------------------------------------------


def test_xlsx_com_data_em_branco_ou_sem_a_coluna_usa_o_fim_do_treino(cliente):
    em_branco = _planilha([['item A', 'diaria', None, 1]])
    sem_coluna = _planilha(
        [['item A', 'semanal']], cabecalho=('sku', 'granularidade')
    )

    primeira = _enviar(cliente, em_branco)
    segunda = _enviar(cliente, sem_coluna)

    assert primeira.status_code == segunda.status_code == HTTP_OK
    ref = lambda r: r.json()['resultados'][0]['previsao']['data_referencia']  # noqa: E731
    assert ref(primeira) == str(FIM_DIARIA)
    assert ref(segunda) == str(FIM_SEMANAL - timedelta(6))


def test_xlsx_valido_com_datas_do_excel_e_em_texto(cliente):
    conteudo = _planilha([
        ['item A', 'diaria', date(2021, 8, 18), 2],
        ['B', 'semanal', '2021-08-15', None],
        ['item A', 'semanal', date(2021, 8, 15), 2.0],
    ])

    resposta = _enviar(cliente, conteudo)

    assert resposta.status_code == HTTP_OK
    corpo = resposta.json()
    assert [r['linha'] for r in corpo['resultados']] == [2, 3, 4]
    assert [len(r['previsao']['periodos']) for r in corpo['resultados']] == [
        2,
        1,
        2,
    ]


def test_xlsx_aponta_linha_e_coluna_de_cada_erro(cliente):
    antes = _gravadas(cliente)
    conteudo = _planilha([
        ['item A', 'diaria', date(2021, 8, 18), 1],
        ['item Z', 'diaria', date(2021, 8, 18), 1],
        ['item A', 'diaria', date(2021, 8, 18), 0],
        ['item A', 'diaria', 'abc', 1],
        ['item A', None, date(2021, 8, 18), 1],
        ['item A', 'mensal', date(2021, 8, 18), 1],
        ['item A', 'diaria', date(2021, 8, 18), 2.5],
    ])

    resposta = _enviar(cliente, conteudo)

    assert resposta.status_code == HTTP_INVALIDO
    campos = _campos(resposta)
    assert 'SKU desconhecido' in campos['linha 3 / sku']
    assert campos['linha 4 / horizonte'] == 'deve ser maior ou igual a 1'
    assert 'linha 5 / data_referencia' in campos
    assert campos['linha 6 / granularidade'] == 'campo obrigatório ausente'
    assert 'mensal não é servida' in campos['linha 7 / granularidade']
    assert campos['linha 8 / horizonte'] == 'deve ser um número inteiro'
    assert not any(c.startswith('linha 2') for c in campos)
    assert _gravadas(cliente) == antes


def test_xlsx_zero_informado_nao_vira_o_padrao(cliente):
    resposta = _enviar(
        cliente, _planilha([['item A', 'diaria', date(2021, 8, 18), 0]])
    )

    assert resposta.status_code == HTTP_INVALIDO
    assert 'linha 2 / horizonte' in _campos(resposta)


def test_xlsx_erros_de_regra_de_negocio_apontam_a_linha(cliente):
    conteudo = _planilha([
        ['item C', 'diaria', date(2021, 8, 18), 1],
        ['item A', 'diaria', date(2021, 8, 18), 15],
        ['item A', 'diaria', date(2021, 8, 1), 1],
    ])

    campos = _campos(_enviar(cliente, conteudo))

    assert 'Sem modelo para item C/diaria' in campos['linha 2']
    assert 'Horizonte 15 acima do limite' in campos['linha 3 / horizonte']
    assert 'anterior ao fim do treino' in campos['linha 4 / data_referencia']


def test_xlsx_ignora_linhas_em_branco_e_preserva_a_numeracao(cliente):
    conteudo = _planilha([
        [None, None, None, None],
        ['item A', 'diaria', date(2021, 8, 18), 1],
        [None, None, None, None],
        ['item Z', 'diaria', date(2021, 8, 18), 1],
    ])

    campos = _campos(_enviar(cliente, conteudo))

    assert list(campos) == ['linha 5 / sku']


def test_xlsx_sem_coluna_obrigatoria_aponta_o_cabecalho(cliente):
    conteudo = _planilha([['item A', 1]], cabecalho=('sku', 'horizonte'))

    campos = _campos(_enviar(cliente, conteudo))

    assert list(campos) == ['linha 1 / granularidade']


def test_xlsx_cabecalho_aceita_caixa_e_ordem_diferentes(cliente):
    conteudo = _planilha(
        [[date(2021, 8, 18), 'item A', 'diaria']],
        cabecalho=('Data_Referencia', 'SKU', 'Granularidade'),
    )

    resposta = _enviar(cliente, conteudo)

    assert resposta.status_code == HTTP_OK


def test_arquivo_que_nao_e_xlsx_e_rejeitado(cliente):
    resposta = _enviar(cliente, b'isto nao e uma planilha')

    assert resposta.status_code == HTTP_INVALIDO
    assert list(_campos(resposta)) == ['arquivo']


def test_planilha_sem_solicitacoes_e_rejeitada(cliente):
    resposta = _enviar(cliente, _planilha([]))

    assert resposta.status_code == HTTP_INVALIDO
    assert list(_campos(resposta)) == ['arquivo']


def test_planilha_grande_demais_e_rejeitada(cliente):
    linhas = [['item A', 'diaria', date(2021, 8, 18), 1]] * (LIMITE_LOTE + 1)

    resposta = _enviar(cliente, _planilha(linhas))

    assert resposta.status_code == HTTP_INVALIDO
    assert f'linha {LIMITE_LOTE + 2}' in _campos(resposta)


def test_rotas_estao_no_swagger(cliente):
    esquema = cliente.get('/openapi.json').json()

    for caminho, metodo in (
        ('/predicoes/lote', 'post'),
        ('/predicoes/lote/xlsx', 'post'),
        ('/predicoes/lote/template.xlsx', 'get'),
    ):
        rota = esquema['paths'][caminho][metodo]
        assert rota['summary']
        assert rota['description'] if metodo == 'post' else True
    assert '422' in esquema['paths']['/predicoes/lote']['post']['responses']
